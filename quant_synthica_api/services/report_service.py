"""Report generation service for institutional Excel workbooks and HTML tearsheets."""

import io
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from quant_synthica_api.services.symbol_resolver import SymbolResolver
from quant_synthica_api.services.market_service import MarketService
from quant_synthica_api.services.fundamentals_service import FundamentalsService
from quant_synthica_api.services.valuation_service import ValuationService
from quant_synthica_api.quant.engine import QuantEngine


class ReportService:
    def __init__(
        self,
        market_service: MarketService,
        fundamentals_service: FundamentalsService,
        valuation_service: ValuationService,
    ):
        self.market = market_service
        self.fundamentals = fundamentals_service
        self.valuation = valuation_service

    def generate_excel_workbook(self, symbol_input: str) -> bytes:
        """Generate institutional multi-tab Excel financial workbook."""
        resolved = SymbolResolver.resolve(symbol_input)
        quote = self.market.get_quote(resolved.canonical).data
        hist = self.market.get_history(resolved.canonical, period="1y", interval="1d").data
        val = self.valuation.get_valuation_overview(resolved.canonical)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style templates
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # 1. Sheet: Overview
        ws_overview = wb.create_sheet(title="Company Overview")
        ws_overview["A1"] = f"{resolved.canonical} - Institutional Factsheet"
        ws_overview["A1"].font = title_font

        pb_ratio = None
        try:
            fund_res = self.fundamentals.get_unified_fundamentals(resolved.canonical)
            pb_ratio = fund_res.data.pb_ratio
        except Exception:
            pass

        overview_data = [
            ("Symbol", resolved.canonical),
            ("Current Price", quote.price),
            ("Day Change", quote.change),
            ("Day Change %", f"{quote.change_percent}%" if quote.change_percent else "-"),
            ("Volume", quote.volume),
            ("Market Cap", quote.market_cap),
            ("52 Week High", quote.fifty_two_week_high),
            ("52 Week Low", quote.fifty_two_week_low),
            ("PE Ratio", quote.pe_ratio),
            ("PB Ratio", pb_ratio),
            ("Dividend Yield", f"{quote.dividend_yield}%" if quote.dividend_yield else "-"),
        ]

        ws_overview["A3"] = "Metric"
        ws_overview["B3"] = "Value"
        ws_overview["A3"].fill = ws_overview["B3"].fill = header_fill
        ws_overview["A3"].font = ws_overview["B3"].font = header_font

        for row_idx, (k, v) in enumerate(overview_data, start=4):
            ws_overview[f"A{row_idx}"] = k
            ws_overview[f"A{row_idx}"].font = bold_font
            ws_overview[f"A{row_idx}"].border = thin_border
            ws_overview[f"B{row_idx}"] = v
            ws_overview[f"B{row_idx}"].font = regular_font
            ws_overview[f"B{row_idx}"].border = thin_border

        # 2. Sheet: Valuation Models
        ws_val = wb.create_sheet(title="Valuation Models")
        ws_val["A1"] = "Quantitative Valuation & Health Scores"
        ws_val["A1"].font = title_font

        val_rows = [
            ("DCF Fair Value", val["dcf_model"].get("fair_value_per_share")),
            ("DCF Margin of Safety %", val["dcf_model"].get("margin_of_safety_pct")),
            ("DCF Undervalued?", val["dcf_model"].get("is_undervalued")),
            ("Piotroski F-Score", f"{val['piotroski_f_score'].get('f_score')} / 9"),
            ("Piotroski Rating", val["piotroski_f_score"].get("rating")),
            ("Altman Z-Score", val["altman_z_score"].get("z_score")),
            ("Altman Zone", val["altman_z_score"].get("zone")),
            ("Altman Bankruptcy Risk", val["altman_z_score"].get("bankruptcy_risk")),
            ("Graham Number", val["graham_valuation"].get("graham_number")),
            ("Graham Undervalued?", val["graham_valuation"].get("is_undervalued_graham")),
        ]

        ws_val["A3"] = "Valuation Dimension"
        ws_val["B3"] = "Assessment"
        ws_val["A3"].fill = ws_val["B3"].fill = header_fill
        ws_val["A3"].font = ws_val["B3"].font = header_font

        for row_idx, (k, v) in enumerate(val_rows, start=4):
            ws_val[f"A{row_idx}"] = k
            ws_val[f"A{row_idx}"].font = bold_font
            ws_val[f"A{row_idx}"].border = thin_border
            ws_val[f"B{row_idx}"] = str(v) if v is not None else "-"
            ws_val[f"B{row_idx}"].font = regular_font
            ws_val[f"B{row_idx}"].border = thin_border

        # 3. Sheet: Historical Prices
        ws_hist = wb.create_sheet(title="Historical Prices")
        hist_headers = ["Date", "Open", "High", "Low", "Close", "Volume"]
        for col_idx, h in enumerate(hist_headers, start=1):
            cell = ws_hist.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, bar in enumerate(hist, start=2):
            ws_hist.cell(row=row_idx, column=1, value=bar.date).font = regular_font
            ws_hist.cell(row=row_idx, column=2, value=bar.open).font = regular_font
            ws_hist.cell(row=row_idx, column=3, value=bar.high).font = regular_font
            ws_hist.cell(row=row_idx, column=4, value=bar.low).font = regular_font
            ws_hist.cell(row=row_idx, column=5, value=bar.close).font = regular_font
            ws_hist.cell(row=row_idx, column=6, value=bar.volume).font = regular_font

        # Auto-adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def generate_tearsheet_html(self, symbol_input: str) -> str:
        """Generate institutional HTML printable tearsheet."""
        resolved = SymbolResolver.resolve(symbol_input)
        quote = self.market.get_quote(resolved.canonical).data
        val = self.valuation.get_valuation_overview(resolved.canonical)
        hist = self.market.get_history(resolved.canonical, period="1y", interval="1d").data

        quant_summary = QuantEngine.compute_summary(hist, resolved.canonical) if len(hist) > 30 else None
        risk = quant_summary.risk.model_dump() if quant_summary else {}
        tech = quant_summary.technical.model_dump() if quant_summary else {}


        dcf = val["dcf_model"]
        piot = val["piotroski_f_score"]
        altman = val["altman_z_score"]
        graham = val["graham_valuation"]

        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>QuantSynthica Tearsheet - {resolved.canonical}</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; background: #0f172a; color: #f8fafc; margin: 0; padding: 24px; }}
  .container {{ max-width: 900px; margin: 0 auto; background: #1e293b; border-radius: 12px; padding: 32px; border: 1px solid #334155; box-shadow: 0 10px 25px rgba(0,0,0,0.5); }}
  .header {{ display: flex; justify-content: space-between; align-items: baseline; border-bottom: 2px solid #3b82f6; padding-bottom: 16px; margin-bottom: 24px; }}
  .title {{ font-size: 28px; font-weight: 800; color: #60a5fa; margin: 0; }}
  .badge {{ background: #3b82f6; color: white; padding: 4px 12px; border-radius: 9999px; font-size: 14px; font-weight: bold; }}
  .grid {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 16px; margin-bottom: 24px; }}
  .card {{ background: #0f172a; padding: 16px; border-radius: 8px; border: 1px solid #334155; }}
  .card-label {{ font-size: 12px; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.05em; }}
  .card-value {{ font-size: 20px; font-weight: bold; color: #f8fafc; margin-top: 6px; }}
  .section-title {{ font-size: 18px; font-weight: 700; color: #93c5fd; margin: 24px 0 12px; border-left: 4px solid #3b82f6; padding-left: 10px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 8px; font-size: 14px; }}
  th {{ background: #0f172a; text-align: left; padding: 10px; color: #94a3b8; border-bottom: 1px solid #334155; }}
  td {{ padding: 10px; border-bottom: 1px solid #334155; }}
  .positive {{ color: #4ade80; }}
  .negative {{ color: #f87171; }}
  .footer {{ margin-top: 32px; font-size: 12px; color: #64748b; text-align: center; border-top: 1px solid #334155; padding-top: 16px; }}
</style>
</head>
<body>
<div class="container">
  <div class="header">
    <div>
      <h1 class="title">{resolved.canonical}</h1>
      <span style="color: #94a3b8; font-size: 14px;">Institutional Tearsheet & Quantitative Factsheet</span>
    </div>
    <span class="badge">QuantSynthica API</span>
  </div>

  <div class="grid">
    <div class="card">
      <div class="card-label">Current Price</div>
      <div class="card-value">{quote.price or '-'}</div>
    </div>
    <div class="card">
      <div class="card-label">Change</div>
      <div class="card-value {'positive' if (quote.change or 0) >= 0 else 'negative'}">{quote.change or 0:+.2f} ({quote.change_percent or 0:+.2f}%)</div>
    </div>
    <div class="card">
      <div class="card-label">Market Cap</div>
      <div class="card-value">{f"{quote.market_cap:,.0f}" if quote.market_cap else '-'}</div>
    </div>
    <div class="card">
      <div class="card-label">P/E Ratio</div>
      <div class="card-value">{quote.pe_ratio or '-'}</div>
    </div>
  </div>

  <div class="section-title">Valuation & Solvency Overview</div>
  <table>
    <tr><th>Model / Test</th><th>Metric Value</th><th>Status / Zone</th></tr>
    <tr>
      <td><strong>DCF Valuation</strong></td>
      <td>Fair Value: {dcf.get('fair_value_per_share', '-')}</td>
      <td><span class="{'positive' if dcf.get('is_undervalued') else 'negative'}">Margin of Safety: {dcf.get('margin_of_safety_pct', '-')}%</span></td>
    </tr>
    <tr>
      <td><strong>Piotroski F-Score</strong></td>
      <td>{piot.get('f_score')}/9</td>
      <td>{piot.get('rating')}</td>
    </tr>
    <tr>
      <td><strong>Altman Z-Score</strong></td>
      <td>{altman.get('z_score')}</td>
      <td>{altman.get('zone')} ({altman.get('bankruptcy_risk')})</td>
    </tr>
    <tr>
      <td><strong>Graham Number</strong></td>
      <td>{graham.get('graham_number', '-')}</td>
      <td>{f"Undervalued: {graham.get('is_undervalued_graham')}" if graham.get('graham_number') else 'N/A'}</td>
    </tr>
  </table>

  <div class="section-title">Risk & Volatility Analytics</div>
  <div class="grid">
    <div class="card">
      <div class="card-label">Annualized Volatility</div>
      <div class="card-value">{f"{risk.get('annualized_volatility', 0)*100:.2f}%" if risk else '-'}</div>
    </div>
    <div class="card">
      <div class="card-label">Sharpe Ratio</div>
      <div class="card-value">{f"{risk.get('sharpe_ratio', 0):.2f}" if risk else '-'}</div>
    </div>
    <div class="card">
      <div class="card-label">Max Drawdown</div>
      <div class="card-value negative">{f"{risk.get('max_drawdown', 0)*100:.2f}%" if risk else '-'}</div>
    </div>
    <div class="card">
      <div class="card-label">RSI (14-Day)</div>
      <div class="card-value">{f"{tech.get('rsi_14', 0):.1f}" if tech else '-'}</div>
    </div>
  </div>

  <div class="footer">
    Generated automatically by QuantSynthica Market API • Data aggregated across yfinance, TradingView, and Screener.in.
  </div>
</div>
</body>
</html>
"""
        return html
